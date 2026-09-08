"""Seed a demo "Acme Consulting" workspace through the normal services.

Run from the backend directory:

    uv run python -m scripts.seed_demo

Requires the schema to be migrated (``uv run alembic upgrade head``) and Redis
up so the worker can process the seeded documents (``docker compose up -d``).
The first user created becomes the root system admin (the Acme owner), so the
demo can also walk through the admin console.

Re-running without ``--force`` fails fast if "Acme Consulting" already exists.
"""

from __future__ import annotations

import argparse
import asyncio
import io
from collections.abc import Sequence

import docx as docx_lib
import pymupdf
from arq import create_pool
from arq.connections import RedisSettings
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.config import get_settings
from app.core.security import hash_password
from app.infrastructure.storage import get_storage
from app.models.user import User
from app.models.workspace import Role, Workspace
from app.services import admin_workspaces as admin_ws_service
from app.services import documents as documents_service
from app.services import folders as folders_service
from app.services import tags as tags_service
from app.services.auth import register_user
from app.services.workspaces import create_workspace

DEMO_PASSWORD = "Acme!demo2026"

USERS: list[tuple[str, str, str, Role]] = [
    ("Alex Acme", "alex@acme.example", "Acme owner", Role.owner),
    ("Casey Admin", "casey@acme.example", "Operations lead", Role.admin),
    ("Mara Manager", "mara@acme.example", "People manager", Role.manager),
    ("Sam Rivera", "sam@acme.example", "Software engineer", Role.employee),
    ("Robin Stone", "robin@acme.example", "Designer", Role.employee),
]

FOLDERS = ["Onboarding", "Policies", "Reports", "Projects"]
TAGS = [
    ("HR", "#f59e0b"),
    ("Finance", "#10b981"),
    ("Urgent", "#ef4444"),
    ("Inbox", "#6366f1"),
]


def _make_pdf(filename: str, pages: Sequence[Sequence[str]]) -> bytes:
    document = pymupdf.open()
    for lines in pages:
        page = document.new_page()
        y = 72
        for line in lines:
            page.insert_text((72, y), line, fontsize=12)
            y += 20
    data = document.tobytes()
    document.close()
    return data


def _make_docx(lines: Sequence[str]) -> bytes:
    document = docx_lib.Document()
    document.add_heading(lines[0], level=1)
    for line in lines[1:]:
        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample_docs() -> list[tuple[str, bytes, list[str], list[str]]]:
    """Return (filename, contents, folder_names, tag_names)."""
    onboarding = _make_pdf(
        "onboarding-guide.pdf",
        [
            [
                "Acme Consulting — Onboarding Guide",
                "",
                "Welcome to Acme! Your first week ahead:",
                "1. Day one: laptop, access, and a welcome from your manager.",
                "2. Read the handbook and the vacation policy.",
                "3. Set up your email and calendar aliases.",
            ],
            [
                "Onboarding checklist",
                "",
                "Accounts provisioned by operations",
                "Buddy assigned from your team",
                "First 1:1 scheduled with your manager",
            ],
        ],
    )
    vacation = _make_pdf(
        "vacation-policy.pdf",
        [
            [
                "Acme Consulting — Vacation Policy",
                "",
                "Employees earn 25 paid vacation days per calendar year.",
                "Requests must be submitted two weeks in advance.",
                "Approval from your manager is required.",
            ],
            [
                "Accrual",
                "",
                "Days accrue monthly and carry over up to 10 days.",
                "Urgent requests are reviewed within 24 hours.",
            ],
        ],
    )
    report = _make_docx(
        [
            "Q4 2025 Financial Report",
            "Revenue grew 18% year over year to $2.4M.",
            "Operating margin improved to 14%.",
            "- Consulting services: +22%",
            "- Product licenses: +11%",
            "Outlook: hiring two engineers and a designer in Q1.",
        ]
    )
    readme = (
        "# Acme Internal Tools — Project README\n\n"
        "The internal tools monorepo powers Acme's consulting dashboards.\n\n"
        "## Getting started\n\n"
        "- Clone the repo and run `make setup`.\n\n"
        "- Copy `.env.example` to `.env`.\n\n"
        "## Notes\n\n"
        "Deploys happen on Fridays — coordinate with ops first."
    ).encode()
    notes = (
        "# Meeting notes — 2026-09-01\n\n"
        "## Client kickoff: Meridian project\n\n"
        "Next milestone is the data migration prototype.\n\n"
        "- Owner: Sam\n"
        "- Deadline: end of September\n\n"
        "Action: share the draft timeline in the portal."
    ).encode()

    return [
        ("onboarding-guide.pdf", onboarding, ["Onboarding"], ["HR", "Inbox"]),
        ("vacation-policy.pdf", vacation, ["Policies"], ["HR", "Urgent"]),
        ("q4-financial-report.docx", report, ["Reports"], ["Finance"]),
        ("project-readme.md", readme, ["Projects"], ["Inbox"]),
        ("meeting-notes-2026.md", notes, ["Projects"], ["Inbox"]),
    ]


async def _seed(force: bool) -> None:
    settings = get_settings()
    engine = create_async_engine(settings.database_url)
    Session = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async with Session() as db:
        existing = await db.scalar(select(Workspace).where(Workspace.name == "Acme Consulting"))
        if existing is not None:
            if not force:
                raise SystemExit('A workspace named "Acme Consulting" already exists. Pass --force to skip seeding.')
            print("Already seeded (Acme Consulting exists) — nothing to do.")
            return

        # Create users through the normal registration flow. The first one becomes the root admin.
        # Existing users are reused by email so the seed stays idempotent across runs.
        registered: dict[str, User] = {}
        for full_name, email, _title, _role in USERS:
            user = await db.scalar(select(User).where(User.email == email))
            if user is None:
                user, _personal_ws = await register_user(db, email, hash_password(DEMO_PASSWORD), full_name)
                await db.commit()
            registered[email] = user

        owner = registered["alex@acme.example"]
        workspace = await create_workspace(db, "Acme Consulting", owner_id=owner.id)
        await db.commit()
        print(f"Created workspace 'Acme Consulting' (id={workspace.id})")

        for _full_name, email, _title, role in USERS[1:]:
            await admin_ws_service.add_member(db, workspace.id, registered[email].id, role)
        await db.commit()

        folder_ids: dict[str, object] = {}
        for name in FOLDERS:
            folder = await folders_service.create_folder(db, workspace, name)
            folder_ids[name] = folder.id
        print(f"Created {len(FOLDERS)} folders: {', '.join(FOLDERS)}")

        tag_ids: dict[str, object] = {}
        for name, color in TAGS:
            tag = await tags_service.create_tag(db, workspace, owner.id, name=name, color=color, scope="global")
            tag_ids[name] = tag.id
        print(f"Created {len(TAGS)} tags")

        await db.commit()

        arq = await create_pool(RedisSettings.from_dsn(settings.redis_url))
        storage = get_storage()

        for filename, contents, folder_names, tag_names in _sample_docs():
            document = await documents_service.upload_document(
                db,
                storage,
                arq,
                workspace,
                owner,
                filename,
                io.BytesIO(contents),
                is_public=False,
            )
            await documents_service.set_document_folders(db, document, [folder_ids[name] for name in folder_names])
            await documents_service.set_document_tags(db, document, owner.id, [tag_ids[name] for name in tag_names])
            print(f"Uploaded {filename} (id={document.id})")
        await arq.aclose()

        print(
            "\nSeeded. Log in with any member (password: {!r}):\n- {}\n".format(
                DEMO_PASSWORD, "\n- ".join(f"{name} <{email}>" for name, email, *_ in USERS)
            )
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--force",
        action="store_true",
        help="skip seeding if 'Acme Consulting' already exists.",
    )
    args = parser.parse_args()
    try:
        asyncio.run(_seed(args.force))
    except OSError as exc:
        raise SystemExit(f"Seed failed ({exc}). Is Postgres up? Did you run `alembic upgrade head`?") from None


if __name__ == "__main__":
    main()
